# -*- coding: utf-8 -*-
"""生成 BempDiff 产品介绍 Word 文档（推广用，含真实功能截图）。

用法：python scripts/build_promo_docx.py
输出：docs/promo/BempDiff-产品介绍.docx
"""
import datetime
import pathlib

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path("D:/code/otherProjects/18_comparePakage")
IMG = ROOT / "docs/promo/img"
OUT = ROOT / "docs/promo/BempDiff-产品介绍.docx"
LOGO = ROOT / "bempdiff/bempdiff-logo.png"
BANNER = ROOT / "bempdiff/assets/social-preview-1280x640.png"

ACCENT = RGBColor(0x2B, 0x63, 0xEC)
GREY = RGBColor(0x6B, 0x72, 0x80)
DARK = RGBColor(0x1F, 0x29, 0x37)
CN_BODY = "等线"
CN_HEAD = "微软雅黑"


def set_cn(run, name):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def style_doc(doc):
    n = doc.styles["Normal"]
    n.font.name = CN_BODY
    n.font.size = Pt(10.5)
    n.font.color.rgb = DARK
    n.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    n.paragraph_format.line_spacing = 1.4
    n.paragraph_format.space_after = Pt(6)
    for lvl, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        st = doc.styles[lvl]
        st.font.name = CN_HEAD
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT if lvl != "Heading 3" else DARK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)
        st.paragraph_format.space_before = Pt(14 if lvl == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)


def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=6, font=CN_BODY):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    set_cn(r, font)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            head, rest = it
            r = p.add_run(head)
            r.bold = True
            set_cn(r, CN_BODY)
            r2 = p.add_run(rest)
            set_cn(r2, CN_BODY)
        else:
            r = p.add_run(it)
            set_cn(r, CN_BODY)
        p.paragraph_format.space_after = Pt(2)


def page_number_footer(doc, text):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + "  ·  第 ")
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    set_cn(r, CN_BODY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16")
    rpr.append(sz)
    run.append(rpr)
    fld.append(run)
    p._p.append(fld)
    r2 = p.add_run(" 页")
    r2.font.size = Pt(8)
    r2.font.color.rgb = GREY
    set_cn(r2, CN_BODY)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    tcpr.append(sh)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cn(r, CN_HEAD)
        shade(c, "2B63EC")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5)
            set_cn(r, CN_BODY)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(doc, fname, caption, width_cm=15.6):
    path = IMG / fname
    if not path.exists():
        placeholder(doc, caption + "（截图待补：缺 " + fname + "）")
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.size = Pt(8.5)
    r.italic = True
    r.font.color.rgb = GREY
    set_cn(r, CN_BODY)
    cap.paragraph_format.space_after = Pt(10)
    return True


def placeholder(doc, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("【此处插入功能截图：" + caption + "】")
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = GREY
    set_cn(r, CN_BODY)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "dashed")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:color"), "9CA3AF")
        bd.append(e)
    pPr.append(bd)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)


def code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        set_cn(r, "Consolas")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), "F3F4F6")
        pPr.append(sh)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:color"), "2B63EC")
    bd.append(bot)
    pPr.append(bd)
    p.paragraph_format.space_after = Pt(8)


def build():
    version = "v0.1.2026101201"
    today = datetime.date.today().strftime("%Y 年 %m 月")
    doc = Document()
    style_doc(doc)
    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.2)
    s.right_margin = Cm(2.2)
    page_number_footer(doc, "BempDiff 产品介绍")

    # ---------------- 封面 ----------------
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(70)
        p.add_run().add_picture(str(LOGO), width=Cm(4.6))
    para(doc, "BempDiff", 40, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, 2, CN_HEAD)
    para(doc, "软件构建包差异化对比与 AI 变更解读工具", 15, True, DARK,
         WD_ALIGN_PARAGRAPH.CENTER, 4, CN_HEAD)
    para(doc, "war / jar / zip / 文件夹 · 逐层解包 · 反编译级 diff · AI 风险分析",
         10.5, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, 18)
    if BANNER.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(BANNER), width=Cm(15.6))
    para(doc, "产品介绍 · 推广版", 11, True, DARK, WD_ALIGN_PARAGRAPH.CENTER, 2, CN_HEAD)
    para(doc, "版本 %s    编制日期 %s" % (version, today), 9.5, False, GREY,
         WD_ALIGN_PARAGRAPH.CENTER, 0)
    doc.add_page_break()

    # ---------------- 一、产品简介 ----------------
    doc.add_heading("一、产品简介", level=1)
    para(doc, "BempDiff 是一款面向「发布前比对」场景的桌面工具：把生产环境正在跑的构建包和本次要下发的构建包"
              "同时丢进去，它会自动逐层解包、按业务重要度分层展示差异、把 class 反编译成源码做逐行比对，"
              "并可调用大模型输出风险概览与逐文件深度解读，最后导出 Markdown 报告与差异资产包。")
    para(doc, "它解决的是一个非常具体的痛点：交付物比对长期依赖「人工解压 + 肉眼比目录 + 记事本看 diff」，"
              "既慢又容易漏，尤其是嵌套 jar、混淆前的业务类改动、配置文件键值增删这类差异，"
              "靠人力在成千上万个条目里定位，一次发布往往要占用测试和开发各半天以上。")
    para(doc, "BempDiff 把这件事变成一次拖拽：双击安装即用（安装包内嵌精简 JRE，目标机器无需安装 Java），"
              "无需搭建任何服务端，也可以完全离线运行；AI 能力按需开启，自带 API Key 即可，"
              "AI 不可达时自动降级，比对与反编译功能不受影响。")
    bullets(doc, [
        ("一句话定位：", "把「两个包之间到底改了什么、会影响什么」这个问题，从半天缩短到三分钟。"),
        ("形态：", "Windows 桌面应用（NSIS 安装包）+ 命令行 headless（可嵌 CI）+ 本地 Web 服务三种用法同源。"),
        ("架构：", "Electron 桌面壳 + Vue 3 前端 + Java 21 核心比对引擎（sidecar 进程），职责清晰、可独立演进。"),
    ])

    # ---------------- 二、核心功能 ----------------
    doc.add_heading("二、核心功能", level=1)

    doc.add_heading("2.1 分层差异视图：先看该看的", level=2)
    para(doc, "工具不会把上万个条目平铺给你。差异结果按重要度分成三层：L0 包级清单（war 内部的文件级增删改）、"
              "L1 业务码（命中「内部包前缀」的 class 展开到方法级源码差异）、L2 第三方依赖（按 jar 折叠，"
              "只在版本号或 CRC 变化时提示）。左侧树顶部即时给出「修改 / 新增 / 删除 / 未变」计数，"
              "支持模糊搜索、正则过滤、全部展开与折叠。")
    figure(doc, "s1_tree.png", "图 1  差异文件树：L0 包级与 L1 业务码分层展示，嵌套 jar 已递归展开")

    doc.add_heading("2.2 逐层解包：嵌套 war → jar → zip 一律摊平", level=2)
    para(doc, "真实的 Java 交付物往往是 war 里套 jar、jar 里再套 zip 的场景资源。BempDiff 会递归解包到"
              "「原子条目」再比对，不放过埋在第三层里的一个 properties 键值改动。"
              "解包采用内存 / 磁盘双轨管理：小文件常驻内存，超出软上限自动落盘，"
              "单个任务内存上限 256MB、全局上限 768MB，比对十万级条目也不会把机器吃满。")
    bullets(doc, [
        "递归深度与解包线程数均可配置，超大包首次比对有进度提示；",
        "解包过程与差异树、AI 分析解耦，边解包边可浏览已就绪的部分。",
    ])

    doc.add_heading("2.3 反编译级源码比对：class 也能逐行看", level=2)
    para(doc, "选中任意一个变更的 class，左右双栏直接展示反编译后的 Java 源码差异——默认使用 CFR，"
              "CFR 无法处理时自动降级到 javap，遇到 GBK 编码的历史包也有容错。"
              "差异呈现完全对齐主流代码审阅习惯：")
    bullets(doc, [
        ("视图：", "左右分栏（split）与 Git 风格统一对比（unified）一键切换；"),
        ("粒度：", "整行 / 词级 / 字符级三档循环，能精确到「多了一个 + 1」这种改动；"),
        ("降噪：", "全量内容与「仅差异内容」互切，并可折叠远离变化块的未变更行（对标 Beyond Compare）；"),
        ("导航：", "上一处 / 下一处差异跳转，j/k/n/p 等 IDE 快捷键，Ctrl+F 在差异内容内查找；"),
        ("溯源：", "顶部同时给出两侧 CRC 指纹、文件大小与所属分层，避免「看着一样其实不同」。"),
    ])
    figure(doc, "s2_class_split.png", "图 2  class 反编译双栏比对：CFR 还原源码，改动行高亮并给出侧边标记")
    figure(doc, "s4_class_unified.png",
           "图 3  Git 风格统一对比：新旧双行号列 + 字符级高亮（精确到多出的「+ 1」）+ 右侧差异位置热力条")

    doc.add_heading("2.4 前端资源与配置对比", level=2)
    para(doc, "JS / HTML / CSS 这类压缩过的文本，直接 diff 只会看到「整行都变了」。"
              "BempDiff 会先做格式化美化（js-beautify）再比对，压缩成一行的脚本也能精确到某个函数体；"
              "properties / xml / yml 等配置项按键值差异呈现，新增与删除一目了然。"
              "Office 文档（docx / xlsx / pptx）走 Apache POI 抽取文本后比对，交付物里有文档也能看差异。")
    figure(doc, "s5_js_diff.png", "图 4  前端 JS 资源比对：美化后逐函数对比，新增函数清晰可辨")

    doc.add_heading("2.5 两阶段 AI 分析：从「改了什么」到「要注意什么」", level=2)
    para(doc, "这是 BempDiff 与通用 diff 工具拉开差距的地方。工具把差异结构化后交给大模型，分两个阶段输出：")
    bullets(doc, [
        ("阶段 A · 整体概览：", "对本次发布做总体风险判断，给出「哪些改动是破坏性的、哪些需要回归、"
                                 "建议的测试要点」，流式输出到下方控制台，可一键复制 Markdown；"),
        ("阶段 B · 单文件深读：", "对 Top-K 高优文件逐个深度分析，结论直接落在右侧「单文件」面板，"
                                  "与源码比对同屏，看完结论就能看代码；"),
        ("多任务控制台：", "支持并行发起多类别分析（全局汇总 / 破坏性 / 审计），"
                          "思考过程可折叠、可中止、可重启，不阻塞界面操作。"),
    ])
    para(doc, "成本与合规同样做了闸门：发起前预估 token 消耗并给出费用提示，超过阈值先告警；"
              "内置金融合规脱敏，敏感字段在进入提示词前处理；API Key 掩码显示且默认不落盘；"
              "厂商可配（OpenAI 兼容 / Azure / Ollama 本地模型均可），内网环境可用本地模型完全离线。")
    figure(doc, "s7b_ai_console_empty.png",
           "图 5  AI 分析发起：控制台选择分析类别（整体风险 / 破坏性变更 / 影响范围 / 测试要点 / 自定义问题）后一键发起")
    figure(doc, "s7_ai_entry.png",
           "图 6  单文件分析面板：选中文件即给出状态、分层、类型、大小与所用反编译引擎")
    figure(doc, "s7d_ai_answer.png",
           "图 7  AI 分析结果正文：整体风险等级、影响范围、测试主题与逐文件风险初评均由大模型实时生成"
           "（示例为 demo_v1.war ↔ demo_v2.war 的真实调用结果）", width_cm=15.0)

    doc.add_heading("2.6 智能分类：让风险等级直接长在差异树上", level=2)
    para(doc, "一键对全部变更文件打标：风险等级 + 变更类别，结果直接反映在左侧差异树的着色与角标上。"
              "审阅人不需要逐个点开，扫一眼树就知道红色节点要重点看、绿色节点基本可以放行。"
              "接入大模型时由模型结合语义判定；未接入时按内置启发式给出确定性结果（按文件类型与包路径判定），"
              "因此这一能力在任何环境下都可用，不会因为没有网络或密钥而空着。")
    figure(doc, "s9_ai_tree.png", "图 8  「智能分析」一键打标后的差异树：风险等级以颜色区分，支持按高 / 中 / 低与类别过滤")

    doc.add_heading("2.7 交互式审阅体验", level=2)
    bullets(doc, [
        ("BCompare 风格操作：", "目录树折叠/展开、路径导航历史（后退/前进/向上一层）、拖拽即比；"),
        ("文件与目录同一个入口：", "同一对话框自动识别类型，包和已解压目录都能直接比；"),
        ("多标签对比：", "同时打开多组比对，标签页切换，最近比对会话一键恢复路径对；"),
        ("专注模式：", "隐藏左右栏放大比对视野（Esc 退出），适合投屏评审；"),
        ("主题与语言：", "多套主题（默认 / GitHub 亮色 / GitHub 暗色 / One Dark）+ 强调色，中英双语界面；"),
        ("已读标记：", "审阅过的文件可标记，团队协作时不重复看。"),
    ])
    figure(doc, "s3_class_only_diff.png", "图 9  仅差异内容模式：未变更行折叠，改动块连续呈现")
    figure(doc, "40_diff_dark.png", "图 10  暗色主题下的同一份差异：多套主题与强调色可切换，适配长时间审阅与投屏")

    doc.add_heading("2.8 报告与资产导出", level=2)
    para(doc, "比对结果不止于「看完就算」。工具支持：")
    bullets(doc, [
        "生成 Markdown 差异报告（只写差异内容，不夹带无关清单），可直接贴进发布单或知识库；",
        "导出差异 class 资产 zip、反编译源码包，方便直接交给开发确认或做补丁验证；",
        "导出粒度可选全量 / JAR / CLASS，按需只取业务码；",
        "大包异步导出 + 下载管理，不阻塞界面。",
    ])
    figure(doc, "s6_export.png", "图 11  导出入口：报告、差异资产与导出粒度（全量 / JAR / CLASS），大包走异步导出与下载管理")

    doc.add_heading("2.9 命令行与 CI 集成", level=2)
    para(doc, "同一套核心引擎提供 headless 子命令，可以直接嵌进流水线做「发布包门禁」："
              "比对结果非零差异即让构建失败，或定期比对基线包与最新构建，防止夹带改动。")
    code(doc, [
        "BempDiff.exe compare   <old> <new> [--expand-all]",
        "BempDiff.exe report    <old> <new> [--report out.md] [--top-k N]",
        "BempDiff.exe ci        --left <old> --right <new> --report r.md --json j.json --fail-on-diff",
        "BempDiff.exe export    <old> <new> [--out outdir]",
        "BempDiff.exe decompile <old> <new> [--cfr cfr.jar]",
        "BempDiff.exe ai        <old> <new> [--apikey KEY] [--baseurl URL] [--project <工程目录>]",
        "BempDiff.exe folderdiff <leftDir> <rightDir> [--report <md>]",
        "BempDiff.exe server    [--port 18765] [--webroot <dir>]",
    ])
    para(doc, "ci 子命令按结果返回退出码（0 无差异 / 1 有差异 / 2 参数或执行错误），"
              "配合 --json 输出可被任意流水线解析。", space_after=10)

    # ---------------- 三、应用场景 ----------------
    doc.add_heading("三、典型应用场景", level=1)
    table(doc,
          ["场景", "过去怎么做", "用 BempDiff 之后"],
          [["版本发布前评审",
            "开发口头描述改动范围，测试凭经验决定回归面",
            "直接给出真实差异清单与 AI 风险提示，回归范围有据可依"],
           ["外包 / 分支交付验收",
            "担心「夹带改动」，只能抽查几个文件",
            "全量逐条目比对 + 反编译级 diff，夹带无处可藏"],
           ["生产事故回溯",
            "两个环境包不同步，靠日志猜是哪次发布引入",
            "拉两份包一比即知差异，配合报告留档可追溯"],
           ["依赖与合规审计",
            "第三方 jar 变更靠人工看版本清单",
            "L2 依赖按 jar 折叠提示，CRC 变化即暴露，含二进制资源"],
           ["流水线发布门禁",
            "无自动化手段，或需自建脚本",
            "ci 子命令 + 退出码，零配置接入现有 CI"]],
          widths=[3.6, 5.6, 6.2])

    # ---------------- 四、为什么选 ----------------
    doc.add_heading("四、与常见做法的对比", level=1)
    table(doc,
          ["维度", "人工解压比对", "通用文本 diff 工具", "BempDiff"],
          [["嵌套包处理", "手动反复解压，极易漏层", "不支持，只能看到二进制乱码",
            "自动逐层解包到原子条目"],
           ["class 差异", "反编译工具单看，无法对比", "不支持",
            "CFR 反编译后双栏逐行比对"],
           ["压缩 JS / 资源", "基本放弃", "整行差异不可读", "美化后比对，精确到函数"],
           ["差异降噪", "无", "全量输出", "L0/L1/L2 分层 + 内部包前缀 + 忽略规则"],
           ["风险判断", "凭经验", "无", "两阶段 AI 分析 + 智能分类着色"],
           ["交付物", "口头 / 截图", "纯文本 diff", "Markdown 报告 + 差异资产包"],
           ["落地成本", "低但耗时", "低", "双击安装即用，内嵌 JRE，可离线"]],
          widths=[2.6, 3.9, 4.2, 4.7])

    # ---------------- 五、规格 ----------------
    doc.add_heading("五、技术规格", level=1)
    table(doc, ["项目", "说明"],
          [["支持输入", "war / jar / zip 等构建包（含嵌套包）、已解压目录、文件夹对"],
           ["运行环境", "Windows 10 / 11 x64；安装包内嵌精简 JRE，无需另装 Java；"
                       "其它平台可用 server 子命令 + 浏览器访问"],
           ["技术栈", "Electron（桌面壳）+ Vue 3 / Bootstrap 5 / Vite（前端）+ Java 21 核心引擎（sidecar）"],
           ["反编译引擎", "CFR（主）· javap（降级）· js-beautify（前端资源）"],
           ["AI 接入", "OpenAI 兼容 / Azure / Ollama 等自定义 BaseURL，自带 Key；成本闸门、脱敏、"
                       "不可达时优雅降级"],
           ["资源控制", "内存 / 磁盘双轨；单任务 256MB、全局 768MB 软上限，解包线程与深度可配"],
           ["离线能力", "比对、反编译、报告导出全部离线可用；AI 部分在无网络且无本地模型时降级"],
           ["许可", "MIT License"]],
          widths=[3.2, 12.2])

    # ---------------- 六、快速上手 ----------------
    doc.add_heading("六、快速上手（三步）", level=1)
    bullets(doc, [
        ("第 1 步 · 安装：", "下载 BempDiff-setup.exe，双击安装（内嵌 JRE，目标机无需装 Java）。"),
        ("第 2 步 · 比对：", "把「老包（生产）」与「新包（下发）」拖进顶部两个输入框——文件与目录同一入口，"
                            "类型自动识别；点击「开始比对」，等待差异树生成。"),
        ("第 3 步 · 审阅与交付：", "左侧树点选文件 → 中间双栏看源码级差异 → 需要时点工具栏「智能分析」"
                                  "发起 AI 风险分析（下方控制台流式输出）→ 从导出菜单生成报告或差异资产包。"),
    ])
    para(doc, "想让 AI 出结论，只需在 ⚙ 配置中心填厂商 BaseURL 与 API Key（掩码显示、默认不落盘），"
              "点「连接测试」通过后即可；不填 Key 也能完整使用比对与反编译能力。", space_after=10)
    figure(doc, "s11_config.png",
           "图 12  配置中心：模型与密钥、解析与导出、差异树过滤、界面与高级——所有配置在界面完成，API Key 掩码显示且默认不落盘")

    # ---------------- 七、行动号召 ----------------
    doc.add_heading("七、获取与支持", level=1)
    para(doc, "如果你也在为「这次发布到底改了什么」反复拉人确认，BempDiff 值得在你的下一版本发布里试一次——"
              "安装包解压即用，第一次比对只需要三分钟。")
    bullets(doc, [
        ("下载：", "GitHub Releases 获取最新 BempDiff-setup.exe（%s）与变更说明。" % version),
        ("文档：", "仓库根 README 提供 GUI / 命令行 / CI 集成的完整用法与架构说明。"),
        ("试用建议：", "拿最近一次已发布的两个包做一次历史比对，用已知的改动清单校验工具的差异结果——"
                      "这是评估工具是否可信最快的方式。"),
    ])
    para(doc, "项目地址：https://github.com/wangnan05563/bempDiff", 10, True, ACCENT, space_after=2)
    para(doc, "许可：MIT　·　形态：桌面应用 / 命令行 / 本地 Web 服务", 9, False, GREY)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print("saved:", build())
